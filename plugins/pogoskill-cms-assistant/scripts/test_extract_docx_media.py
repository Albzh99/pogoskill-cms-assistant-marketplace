import importlib.util
import tempfile
import unittest
from pathlib import Path


SPEC = importlib.util.spec_from_file_location("extract_docx_media", Path(__file__).with_name("extract-docx-media.py"))
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


class ExtractTests(unittest.TestCase):
    def test_preserves_two_occurrences_of_same_image(self):
        from docx import Document
        from docx.shared import Inches
        from PIL import Image
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            image = root / "source.png"
            Image.new("RGB", (20, 12), "blue").save(image)
            doc = Document()
            doc.add_picture(str(image), width=Inches(1))
            doc.add_paragraph("Between images")
            doc.add_picture(str(image), width=Inches(1))
            source = root / "article.docx"
            doc.save(source)
            result = MODULE.extract(source, root / "out")
            self.assertEqual(2, result["image_occurrences"])
            self.assertEqual(result["images"][0]["source_sha256"], result["images"][1]["source_sha256"])
            self.assertTrue(Path(result["images"][0]["extracted_path"]).is_file())


if __name__ == "__main__":
    unittest.main()
